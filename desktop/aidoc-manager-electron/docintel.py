"""证件资料识别引擎（桌面本地版）。

移植自网页参考版 DocumentIntelligenceService 的核心思路：
- 取证据：有文字层的 PDF / Office / 文本直接抠字；扫描件、图片渲染成页图
- 判类型 + 抠关键字段（规则层，廉价、可离线兜底）
- 主力识别交给 AI（文本模型或视觉模型，提示词见 EXTRACTION_INSTRUCTION），返回结构化字段
- 归一化到统一 schema，进「待确认」

本模块只做「文件 -> 证据 -> 字段」的纯逻辑，AI 远程调用由 backend.py 走服务器（按页计费）。
"""

from __future__ import annotations

import io
import json
import re
from pathlib import Path
from typing import Any

# ───────────────────────── 证件类型与字段 schema ─────────────────────────

DOCUMENT_TYPES = {
    "business_license": "营业执照",
    "trademark_certificate": "商标证",
    "authorization_letter": "授权书",
    "barcode_certificate": "条码证",
    "quality_report": "质检/检测报告",
    "ccc_certificate": "3C 证书",
    "production_license": "生产许可证",
    "hygiene_license": "卫生许可证",
    "product_filing": "产品备案/注册",
    "food_license": "食品许可证",
    "contract": "合同/协议",
    "tax_certificate": "税务/开户资料",
    "other": "其他资料",
}

# 计算「缺失资料」的基线证件类型
REQUIRED_TYPES = [
    "business_license",
    "trademark_certificate",
    "authorization_letter",
    "quality_report",
]

# 主列表「总预览」列（来自 版面样式.xlsx 总预览页）
LIST_COLUMNS = [
    {"key": "company_name", "label": "归属公司"},
    {"key": "document_type_label", "label": "类型"},
    {"key": "issued_at", "label": "有效期起"},
    {"key": "expires_at", "label": "有效期止"},
]

# 顶层（公共）字段——列存储，总预览/筛选/有效期都用。
TOP_FIELDS = [
    "document_type", "company_name", "brand", "certificate_no",
    "issuer", "issued_at", "expires_at", "applicable_scope",
    "ai_summary", "ai_confidence",
]

# 每种类型的有序展示/编辑字段，严格对齐「版面样式.xlsx」。
# source 取值：顶层字段名，或 "extra:<key>"（存进 extra_fields）。
FIELD_PROFILES: dict[str, list[dict[str, str]]] = {
    "business_license": [
        {"label": "公司名称", "source": "company_name"},
        {"label": "注册金", "source": "extra:registered_capital"},
        {"label": "注册时间", "source": "issued_at"},
        {"label": "到期时间", "source": "expires_at"},
        {"label": "经营范围", "source": "applicable_scope"},
        {"label": "编号/税号", "source": "certificate_no"},
        {"label": "法人", "source": "extra:legal_representative"},
        {"label": "地址", "source": "extra:address"},
    ],
    "contract": [
        {"label": "合同编号", "source": "certificate_no"},
        {"label": "甲方", "source": "company_name"},
        {"label": "乙方", "source": "issuer"},
        {"label": "丙方", "source": "extra:party_c"},
        {"label": "类型", "source": "extra:contract_type"},
        {"label": "内容简述", "source": "applicable_scope"},
        {"label": "有效期起", "source": "issued_at"},
        {"label": "有效期止", "source": "expires_at"},
    ],
    "trademark_certificate": [
        {"label": "商标名称", "source": "brand"},
        {"label": "归属公司", "source": "company_name"},
        {"label": "证书编号", "source": "certificate_no"},
        {"label": "有效期起", "source": "issued_at"},
        {"label": "有效期止", "source": "expires_at"},
    ],
    "authorization_letter": [
        {"label": "证书编号", "source": "certificate_no"},
        {"label": "授权方", "source": "company_name"},
        {"label": "被授权方", "source": "issuer"},
        {"label": "授权品牌", "source": "brand"},
        {"label": "授权内容", "source": "applicable_scope"},
        {"label": "有效期起", "source": "issued_at"},
        {"label": "有效期止", "source": "expires_at"},
    ],
    "quality_report": [
        {"label": "证书编号", "source": "certificate_no"},
        {"label": "归属公司", "source": "company_name"},
        {"label": "检测商品名", "source": "applicable_scope"},
        {"label": "检测机构", "source": "issuer"},
        {"label": "有效期起", "source": "issued_at"},
        {"label": "有效期止", "source": "expires_at"},
    ],
    "ccc_certificate": [
        {"label": "证书编号", "source": "certificate_no"},
        {"label": "归属公司", "source": "company_name"},
        {"label": "商品名", "source": "applicable_scope"},
        {"label": "出具机构", "source": "issuer"},
        {"label": "有效期起", "source": "issued_at"},
        {"label": "有效期止", "source": "expires_at"},
    ],
}

# 版面样式未细化的类型（条码证/卫生许可证/生产许可证/产品备案/食品许可证/税务/其他）走通用字段。
DEFAULT_PROFILE = [
    {"label": "归属公司", "source": "company_name"},
    {"label": "证件编号", "source": "certificate_no"},
    {"label": "适用/范围", "source": "applicable_scope"},
    {"label": "有效期起", "source": "issued_at"},
    {"label": "有效期止", "source": "expires_at"},
]


def field_profile(document_type: str) -> list[dict[str, str]]:
    return FIELD_PROFILES.get(document_type, DEFAULT_PROFILE)


# 专属字段（extra_fields）——只保留 FIELD_PROFILES 里 extra: 的字段，用于归一化时过滤。
EXTRA_FIELD_SCHEMAS: dict[str, list[dict[str, str]]] = {
    "business_license": [
        {"key": "registered_capital", "label": "注册金", "type": "text"},
        {"key": "legal_representative", "label": "法人", "type": "text"},
        {"key": "address", "label": "地址", "type": "text"},
    ],
    "contract": [
        {"key": "party_c", "label": "丙方", "type": "text"},
        {"key": "contract_type", "label": "类型", "type": "text"},
    ],
}

# 规则层类型判断关键词（命中越多越靠前）
TYPE_KEYWORDS: dict[str, list[str]] = {
    "business_license": ["营业执照", "统一社会信用代码", "经营范围", "法定代表人"],
    "trademark_certificate": ["商标注册证", "商标注册", "注册商标", "核定使用商品", "商标局"],
    "authorization_letter": ["授权书", "授权委托", "兹授权", "授权方", "被授权"],
    "barcode_certificate": ["中国商品条码", "厂商识别代码", "条码", "物编注字", "系统成员证书"],
    "quality_report": ["检验报告", "检测报告", "质检报告", "报告编号", "检验结论", "委托单位"],
    "ccc_certificate": ["强制性产品认证", "3C", "CCC", "中国国家强制性"],
    "production_license": ["生产许可证", "全国工业产品生产许可证", "许可证编号"],
    "hygiene_license": ["卫生许可证", "消毒产品", "卫消证字", "生产企业卫生"],
    "product_filing": ["产品备案", "备案凭证", "注册证", "备案号"],
    "food_license": ["食品经营许可证", "食品生产许可证", "SC", "食品许可"],
    "contract": ["合同", "协议", "甲方", "乙方", "采购"],
    "tax_certificate": ["开户许可证", "纳税人", "银行账号", "税务登记"],
}


def extraction_instruction() -> str:
    """识别提示词（字段对齐「版面样式.xlsx」）。"""
    types = "，".join(f"{k}={v}" for k, v in DOCUMENT_TYPES.items())
    return (
        "你是供应商资料库的文档管理员。请判断证件类型并提取结构化字段，只输出 JSON，不要解释、不要 Markdown。\n\n"
        f"可选资料类型(document_type 取等号左边的代码)：{types}\n\n"
        "统一顶层字段（所有类型都尽量填）：\n"
        "- company_name 归属公司：营业执照填公司名称；合同填甲方；商标证/质检/3C 填归属公司/注册人；授权书填授权方。\n"
        "- certificate_no 证件编号：营业执照填统一社会信用代码/税号；合同填合同编号；商标证/授权书/质检/3C 填证书编号。\n"
        "- brand 品牌：商标证填商标名称；授权书填授权品牌；其它无则留空。\n"
        "- issuer：合同填乙方；授权书填被授权方；质检报告填检测机构；3C 填出具机构。\n"
        "- applicable_scope：营业执照填经营范围；合同填内容简述；授权书填授权内容；质检报告填检测商品名；3C 填商品名。\n"
        "- issued_at 有效期起 / expires_at 有效期止：营业执照=注册时间/到期时间；其它=证件有效期起止；授权书=授权期限起止；质检报告若只有报告日期则填 issued_at。\n\n"
        "各类型专属字段放进 extra_fields（对象，只填下面列出的 key，没有就不填）：\n"
        "- business_license 营业执照：registered_capital(注册金)、legal_representative(法人)、address(地址)。\n"
        "- contract 合同：party_c(丙方)、contract_type(合同类型，如采购/服务/代理)。\n"
        "- 其它类型 extra_fields 留空对象 {}。\n\n"
        "日期一律 YYYY-MM-DD。识别不出的字段填空字符串。ai_confidence 是 0-100 的整数。\n"
        "返回 JSON：{document_type, company_name, brand, certificate_no, issuer, issued_at, expires_at, applicable_scope, extra_fields, tags, ai_summary, ai_confidence}"
    )


# ───────────────────────── 取证据：文字 / 页图 ─────────────────────────

IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff"}


def extract_text(path: Path) -> str:
    """尽量本地抠出文字层；扫描件会返回很短/空字符串。"""
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            import fitz  # PyMuPDF
            doc = fitz.open(str(path))
            text = "".join(page.get_text() for page in doc)
            doc.close()
            return text.strip()
        if suffix == ".docx":
            from docx import Document
            document = Document(str(path))
            return "\n".join(p.text for p in document.paragraphs).strip()
        if suffix in {".txt", ".csv"}:
            return path.read_text("utf-8", errors="ignore").strip()
        if suffix in {".xlsx", ".xlsm"}:
            import openpyxl
            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            parts: list[str] = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    parts.append(" ".join(str(c) for c in row if c is not None))
            wb.close()
            return "\n".join(parts).strip()
    except Exception:
        return ""
    return ""


def render_pages_png(path: Path, max_pages: int = 2, dpi: int = 150) -> list[bytes]:
    """把 PDF 前几页或图片渲染成 PNG 字节，喂视觉 AI。"""
    suffix = path.suffix.lower()
    out: list[bytes] = []
    try:
        if suffix == ".pdf":
            import fitz
            doc = fitz.open(str(path))
            zoom = dpi / 72.0
            matrix = fitz.Matrix(zoom, zoom)
            for page in doc[:max_pages]:
                out.append(page.get_pixmap(matrix=matrix).tobytes("png"))
            doc.close()
        elif suffix in IMAGE_EXTS:
            from PIL import Image
            with Image.open(path) as im:
                im = im.convert("RGB")
                # 控制尺寸，长边压到 1600，省流量
                long_side = max(im.size)
                if long_side > 1600:
                    scale = 1600 / long_side
                    im = im.resize((round(im.width * scale), round(im.height * scale)))
                buf = io.BytesIO()
                im.save(buf, "PNG")
                out.append(buf.getvalue())
    except Exception:
        return []
    return out


def needs_vision(path: Path, text: str) -> bool:
    """图片必走视觉；PDF 文字层太少（扫描件）也走视觉。"""
    suffix = path.suffix.lower()
    if suffix in IMAGE_EXTS:
        return True
    if suffix == ".pdf":
        return len(text) < 30
    return False


# ───────────────────────── 规则层：类型 / 日期 / 证件号 ─────────────────────────

def detect_type(text: str) -> str:
    if not text:
        return "other"
    scores: dict[str, int] = {}
    for type_key, words in TYPE_KEYWORDS.items():
        score = sum(text.count(w) for w in words)
        if score:
            scores[type_key] = score
    if not scores:
        return "other"
    return max(scores, key=scores.get)


_DATE_PATTERNS = [
    re.compile(r"(\d{4})\s*[-/年.]\s*(\d{1,2})\s*[-/月.]\s*(\d{1,2})"),
    re.compile(r"(\d{4})\s*[-/年.]\s*(\d{1,2})\s*月"),
]


def extract_dates(text: str) -> list[str]:
    found: list[str] = []
    for pat in _DATE_PATTERNS:
        for m in pat.finditer(text):
            y = int(m.group(1))
            mo = int(m.group(2))
            d = int(m.group(3)) if m.lastindex and m.lastindex >= 3 else 1
            if 1900 <= y <= 2100 and 1 <= mo <= 12 and 1 <= d <= 31:
                iso = f"{y:04d}-{mo:02d}-{d:02d}"
                if iso not in found:
                    found.append(iso)
    return found


_CERT_NO_LABELS = ["证书编号", "报告编号", "证号", "编号", "注册号", "许可证编号", "卫消证字", "备案号"]


def extract_cert_no(text: str) -> str:
    for label in _CERT_NO_LABELS:
        m = re.search(label + r"[:：]?\s*([A-Za-z0-9（）()\-/．.]{4,40})", text)
        if m:
            return m.group(1).strip(" .，,；;")
    return ""


def extract_by_labels(text: str, labels: list[str], limit: int = 180) -> str:
    for label in labels:
        m = re.search(re.escape(label) + r"[:：]?\s*([^\n\r]{2," + str(limit) + r"})", text)
        if m:
            return m.group(1).strip(" .，,；;：:")
    return ""


def normalize_date(value: Any) -> str:
    """各种日期写法 → YYYY-MM-DD；认不出返回空。"""
    text = str(value or "").strip()
    if not text:
        return ""
    m = re.search(r"(\d{4})\D{0,2}(\d{1,2})\D{0,2}(\d{1,2})", text)
    if m:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if 1900 <= y <= 2100 and 1 <= mo <= 12 and 1 <= d <= 31:
            return f"{y:04d}-{mo:02d}-{d:02d}"
    m = re.search(r"(\d{4})\D{0,2}(\d{1,2})", text)
    if m:
        y, mo = int(m.group(1)), int(m.group(2))
        if 1900 <= y <= 2100 and 1 <= mo <= 12:
            return f"{y:04d}-{mo:02d}-01"
    return ""


def rule_suggestion(text: str) -> dict[str, Any]:
    """纯规则识别：AI 不可用时的兜底，也用于给 AI 结果补缺。"""
    doc_type = detect_type(text)
    dates = extract_dates(text)
    cert_no = extract_cert_no(text)
    company = extract_by_labels(text, ["公司名称", "单位名称", "企业名称", "注册人", "委托单位", "授权方"])
    confidence = 30
    if doc_type != "other":
        confidence += 20
    if cert_no:
        confidence += 10
    if company:
        confidence += 10
    return {
        "document_type": doc_type,
        "company_name": company,
        "brand": "",
        "certificate_no": cert_no,
        "issuer": "",
        "issued_at": dates[0] if dates else "",
        "expires_at": dates[-1] if len(dates) > 1 else "",
        "applicable_scope": "",
        "extra_fields": {},
        "tags": [t for t in [DOCUMENT_TYPES.get(doc_type), company] if t],
        "ai_summary": "",
        "ai_confidence": min(70, confidence),
    }


# ───────────────────────── 归一化 / 合并 ─────────────────────────

def parse_ai_json(content: str) -> dict[str, Any] | None:
    text = (content or "").strip()
    if not text:
        return None
    if text.startswith("```"):
        text = re.sub(r"^```[a-zA-Z]*\s*|\s*```$", "", text).strip()
    try:
        data = json.loads(text)
    except Exception:
        m = re.search(r"\{.*\}", text, re.S)
        if not m:
            return None
        try:
            data = json.loads(m.group(0))
        except Exception:
            return None
    return data if isinstance(data, dict) else None


def normalize_suggestion(raw: dict[str, Any]) -> dict[str, Any]:
    """把 AI/规则的原始结果清洗成统一 schema。"""
    raw = raw or {}
    doc_type = str(raw.get("document_type") or "other")
    if doc_type not in DOCUMENT_TYPES:
        doc_type = "other"

    extra_in = raw.get("extra_fields")
    extra: dict[str, Any] = extra_in if isinstance(extra_in, dict) else {}
    # 只保留该类型 schema 里定义的字段
    allowed = {f["key"] for f in EXTRA_FIELD_SCHEMAS.get(doc_type, [])}
    extra = {k: v for k, v in extra.items() if k in allowed and str(v or "").strip() != ""}

    tags_in = raw.get("tags")
    tags = [str(t).strip() for t in tags_in if str(t or "").strip()] if isinstance(tags_in, list) else []

    try:
        confidence = int(float(raw.get("ai_confidence") or 0))
    except (TypeError, ValueError):
        confidence = 0
    confidence = max(0, min(100, confidence))

    return {
        "document_type": doc_type,
        "document_type_label": DOCUMENT_TYPES.get(doc_type, DOCUMENT_TYPES["other"]),
        "company_name": str(raw.get("company_name") or "").strip(),
        "brand": str(raw.get("brand") or "").strip(),
        "certificate_no": str(raw.get("certificate_no") or "").strip(),
        "issuer": str(raw.get("issuer") or "").strip(),
        "issued_at": normalize_date(raw.get("issued_at")),
        "expires_at": normalize_date(raw.get("expires_at")),
        "applicable_scope": str(raw.get("applicable_scope") or "").strip(),
        "extra_fields": extra,
        "tags": tags[:8],
        "ai_summary": str(raw.get("ai_summary") or "").strip(),
        "ai_confidence": confidence,
    }


def merge_rule_and_ai(rule: dict[str, Any], ai: dict[str, Any] | None) -> dict[str, Any]:
    """以 AI 结果为主，AI 缺的用规则补；都用归一化后的字段。"""
    if not ai:
        return normalize_suggestion(rule)
    merged = dict(ai)
    if merged.get("document_type", "other") == "other" and rule.get("document_type") != "other":
        merged["document_type"] = rule["document_type"]
    for key in ("company_name", "certificate_no", "issued_at", "expires_at", "applicable_scope", "brand", "issuer"):
        if not str(merged.get(key) or "").strip() and str(rule.get(key) or "").strip():
            merged[key] = rule[key]
    if not merged.get("tags") and rule.get("tags"):
        merged["tags"] = rule["tags"]
    return normalize_suggestion(merged)


def project_to_profile(suggestion: dict[str, Any]) -> list[dict[str, str]]:
    """按类型版面（FIELD_PROFILES）把识别结果投影成有序 [{label, source, value}]，给前端渲染/编辑。"""
    extra = suggestion.get("extra_fields") or {}
    rows: list[dict[str, str]] = []
    for field in field_profile(str(suggestion.get("document_type") or "other")):
        source = field["source"]
        if source.startswith("extra:"):
            value = extra.get(source.split(":", 1)[1], "")
        elif source == "document_type_label":
            value = suggestion.get("document_type_label", "")
        else:
            value = suggestion.get(source, "")
        rows.append({"label": field["label"], "source": source, "value": str(value or "")})
    return rows
